"""Excel workbook generator creating 6 professionally styled sheets via openpyxl."""
import io
from typing import Any, Dict, List
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from analysis.model_summary import generate_model_summary
from analysis.jd_summary import generate_jd_summary
from analysis.category_summary import generate_category_summary
from export.csv_export import CSV_28_COLUMNS


def style_header_row(ws, row_idx=1, fill_color="1E293B", font_color="FFFFFF"):
    """Applies clean dark headers with bold white text."""
    fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    font = Font(name="Calibri", size=11, bold=True, color=font_color)
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)

    for cell in ws[row_idx]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[row_idx].height = 28


def autofit_columns(ws, max_len_cap=45):
    """Auto-adjusts column widths with padding and maximum width capping."""
    thin = Side(border_style="thin", color="E2E8F0")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)
    
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            cell.border = border
            if cell.row > 1:
                cell.font = Font(name="Calibri", size=10)
                if isinstance(cell.value, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            if cell.value:
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), max_len_cap)


def export_comparisons_to_xlsx(comparisons: List[Dict[str, Any]]) -> bytes:
    """Creates a 6-sheet styled Excel workbook."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # ----------------------------------------------------
    # Sheet 1: Model Summary (Leaderboard)
    # ----------------------------------------------------
    ws_model = wb.create_sheet(title="Model Summary")
    ws_model.append([
        "Rank", "AI Model", "Status", "Evaluated JDs", "Avg Advantage", "Std Dev",
        "Avg Rel Increase %", "Avg AI Overall", "Avg Human Overall",
        "Avg Keyword", "Avg Skills", "Avg Formatting", "Avg Role", "Avg Impact",
        "AI Wins", "Human Wins", "Ties", "Win Rate %"
    ])
    style_header_row(ws_model, 1, fill_color="0F172A")
    
    model_summaries = generate_model_summary(comparisons)
    for m in model_summaries:
        ws_model.append([
            m.get("rank") or "Provisional",
            m["model"],
            m["status"],
            m["jd_count"],
            m["avg_advantage"],
            m["std_advantage"],
            m["avg_relative_increase"],
            m["avg_ai_overall"],
            m["avg_human_overall"],
            m["avg_keyword_match"],
            m["avg_skills_coverage"],
            m["avg_ats_formatting"],
            m["avg_role_alignment"],
            m["avg_impact_metrics"],
            m["ai_wins"],
            m["human_wins"],
            m["ties"],
            m["win_rate_percent"],
        ])
    autofit_columns(ws_model)

    # ----------------------------------------------------
    # Sheet 2: Master Comparisons (All Records)
    # ----------------------------------------------------
    ws_master = wb.create_sheet(title="Master Comparisons")
    ws_master.append(CSV_28_COLUMNS)
    style_header_row(ws_master, 1, fill_color="1E293B")

    for c in comparisons:
        row_vals = [c.get(col) for col in CSV_28_COLUMNS]
        ws_master.append(row_vals)
    autofit_columns(ws_master)

    # ----------------------------------------------------
    # Sheet 3: JD Breakdown
    # ----------------------------------------------------
    ws_jd = wb.create_sheet(title="JD Breakdown")
    ws_jd.append([
        "Job Description", "Models Evaluated", "Avg AI Score", "Avg Human Score",
        "Avg Advantage", "Std Dev", "AI Wins", "Human Wins", "Ties", "AI Win Rate %"
    ])
    style_header_row(ws_jd, 1, fill_color="334155")
    
    jd_summaries = generate_jd_summary(comparisons)
    for j in jd_summaries:
        ws_jd.append([
            j["jd"],
            j["model_count"],
            j["avg_ai_score"],
            j["avg_human_score"],
            j["avg_advantage"],
            j["std_advantage"],
            j["ai_wins"],
            j["human_wins"],
            j["ties"],
            j["ai_win_rate_percent"],
        ])
    autofit_columns(ws_jd)

    # ----------------------------------------------------
    # Sheet 4: Category Breakdown
    # ----------------------------------------------------
    ws_cat = wb.create_sheet(title="Category Breakdown")
    ws_cat.append([
        "Category", "Weight", "Avg AI Score", "Avg Human Score",
        "Avg Advantage", "Std Dev", "AI Wins", "Human Wins", "Ties"
    ])
    style_header_row(ws_cat, 1, fill_color="475569")
    
    cat_summaries = generate_category_summary(comparisons)
    for cat in cat_summaries:
        ws_cat.append([
            cat["category"],
            cat["weight"],
            cat["avg_ai_score"],
            cat["avg_human_score"],
            cat["avg_advantage"],
            cat["std_advantage"],
            cat["ai_wins"],
            cat["human_wins"],
            cat["ties"],
        ])
    autofit_columns(ws_cat)

    # ----------------------------------------------------
    # Sheet 5: Model x JD Advantage Matrix
    # ----------------------------------------------------
    ws_matrix = wb.create_sheet(title="Advantage Matrix")
    matrix_headers = ["AI Model", "JD1", "JD2", "JD3", "JD4", "JD5", "Average Advantage"]
    ws_matrix.append(matrix_headers)
    style_header_row(ws_matrix, 1, fill_color="1E3A8A")

    comp_lookup = {(c["model"], c["jd"]): c.get("calculated_ats_advantage") for c in comparisons}
    for m in model_summaries:
        model_name = m["model"]
        row_vals = [model_name]
        for jd in ["JD1", "JD2", "JD3", "JD4", "JD5"]:
            val = comp_lookup.get((model_name, jd))
            row_vals.append(val if val is not None else "-")
        row_vals.append(m["avg_advantage"] if m["avg_advantage"] is not None else "-")
        ws_matrix.append(row_vals)
    autofit_columns(ws_matrix)

    # ----------------------------------------------------
    # Sheet 6: Audit & Integrity Log
    # ----------------------------------------------------
    ws_audit = wb.create_sheet(title="Audit Log")
    ws_audit.append([
        "Comparison ID", "Model", "JD", "Version", "Math Consistent",
        "Reported AI", "Calculated AI", "Reported Human", "Calculated Human",
        "Reported Adv", "Calculated Adv", "Claude URL", "Created At"
    ])
    style_header_row(ws_audit, 1, fill_color="047857")

    for c in comparisons:
        ws_audit.append([
            c.get("comparison_id"),
            c.get("model"),
            c.get("jd"),
            c.get("version", 1),
            "YES" if c.get("is_math_consistent", 1) else "DISCREPANCY",
            c.get("ai_overall_reported"),
            c.get("ai_overall_calculated"),
            c.get("human_overall_reported"),
            c.get("human_overall_calculated"),
            c.get("reported_ats_advantage"),
            c.get("calculated_ats_advantage"),
            c.get("claude_chat_url"),
            c.get("created_at"),
        ])
    autofit_columns(ws_audit)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
